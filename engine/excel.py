"""
Genera el Excel formateado de un calendario ya expandido (`schedule.py#subject_to_events`)
— colores por tipo de grupo (Teoría/Prácticas/Trabajo en grupo, con variante
inglés), hoja "Horario" (rejilla semana-a-semana + resumen de horas por
evento) y hoja "Eventos" (lista plana). Heredado de una app anterior,
quitando `append_plan`/`add_plan` (opción de fichero de planificación de temas/profesor —
desactivada en la app original, nunca se ejercitaba en producción, ver
schedule.py).
"""

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter

from schedule import print_subject


def _add_outline_border(sheet, start_row, start_col, end_row, end_col, border_style="thick"):
    """Borde alrededor del PERÍMETRO de un rango (no de cada celda) — se llama varias veces con distintos rangos para dibujar el marco exterior + separadores de cabecera de cada tabla semanal."""
    border_side = Side(style=border_style)
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            cell = sheet.cell(row, col)
            if row == start_row:
                cell.border = Border(top=border_side, right=cell.border.right, bottom=cell.border.bottom, left=cell.border.left)
            if row == end_row:
                cell.border = Border(bottom=border_side, right=cell.border.right, top=cell.border.top, left=cell.border.left)
            if col == start_col:
                cell.border = Border(left=border_side, right=cell.border.right, top=cell.border.top, bottom=cell.border.bottom)
            if col == end_col:
                cell.border = Border(right=border_side, top=cell.border.top, bottom=cell.border.bottom, left=cell.border.left)


def _add_all_borders(sheet, start_row, start_col, end_row, end_col, border_style="thin"):
    border = Border(top=Side(style=border_style), right=Side(style=border_style), bottom=Side(style=border_style), left=Side(style=border_style))
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            sheet.cell(row, col).border = border


def _add_conditional_formatting(sheet, formula, fill_color):
    """Colorea las celdas de la rejilla de horario según el texto del grupo (p.ej. `PL-ENG1` -> naranja) — con formato condicional de Excel, no coloreando celda a celda, así el color se recalcula solo si se edita algo a mano después."""
    fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    dxf = DifferentialStyle(fill=fill)
    rule = Rule(type="expression", dxf=dxf)
    rule.formula = [formula]
    sheet.conditional_formatting.add("A1:Z1000", rule)


def create_subject_excel(subject_events, subject_name, holiday, min_h=None, max_h=None):
    """`subject_events`: eventos ya expandidos de `schedule.py#subject_to_events` (misma entrada que `to_ical`). Devuelve los bytes del .xlsx."""
    data, subject_calendar = print_subject(subject_events, column="group", min_h=min_h, max_h=max_h, weekend=False)

    fmg = 1  # filas de margen al inicio
    cmg = 1  # columnas de margen al inicio
    fst = 1  # filas de separación entre tablas semanales

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        row = fmg
        for lista in data:
            df = pd.DataFrame(lista)
            df.to_excel(writer, startrow=row, startcol=cmg, sheet_name="Horario", index=False, header=False)
            row = row + len(df.index) + fst

        summary = subject_calendar.copy()
        summary["length"] = summary["end"] - summary["start"]
        summary = summary.groupby(["group", "length"]).agg(Horas=("group", "count")).reset_index(names=["Evento", "length"])
        summary["Horas"] = summary["Horas"] * summary["length"]
        summary = summary.loc[summary["Evento"] != holiday][["Evento", "Horas"]]
        summary = summary.groupby("Evento").agg(Horas=("Horas", "sum")).reset_index()
        summary.to_excel(writer, startrow=fmg, startcol=cmg + len(data[0][0]) + 1, sheet_name="Horario", index=False)

    output.seek(0)
    wb = load_workbook(output)
    ws = wb["Horario"]

    start_col = cmg + 1
    end_col = len(data[0][0]) + 4
    for col in range(start_col, end_col + 1):
        for cell in ws[get_column_letter(col)]:
            cell.alignment = Alignment(horizontal="center")

    ws.column_dimensions["B"].width = 3.57
    ws.column_dimensions["C"].width = 3.57

    _add_conditional_formatting(ws, 'A1="A"', "E2EFDA")
    _add_conditional_formatting(ws, 'A1="ENG"', "548235")
    _add_conditional_formatting(ws, 'A1="-"', "C9C9C9")
    _add_conditional_formatting(ws, f'A1="{holiday}"', "C9C9C9")
    _add_conditional_formatting(ws, 'AND(ISNUMBER(SEARCH("PL-ENG", A1)), NOT(ISBLANK(A1)))', "F4B084")
    _add_conditional_formatting(ws, 'AND(ISNUMBER(SEARCH("PL", A1)), NOT(ISBLANK(A1)))', "FCE4D6")
    _add_conditional_formatting(ws, 'AND(ISNUMBER(SEARCH("TG-ENG", A1)), NOT(ISBLANK(A1)))', "FFD966")
    _add_conditional_formatting(ws, 'AND(ISNUMBER(SEARCH("TG", A1)), NOT(ISBLANK(A1)))', "FFF2CC")
    _add_conditional_formatting(ws, 'AND(ISNUMBER(SEARCH("PA-ENG", A1)), NOT(ISBLANK(A1)))', "8EA9DB")
    _add_conditional_formatting(ws, 'AND(ISNUMBER(SEARCH("PA", A1)), NOT(ISBLANK(A1)))', "D9E1F2")

    header_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    for idx in range(len(data)):
        rowno = fmg + len(data[0]) * idx + fst * idx + 1
        ws.merge_cells(start_row=rowno, start_column=cmg + 1, end_row=rowno, end_column=cmg + 2)
        _add_all_borders(ws, rowno, cmg + 1, rowno + len(data[0]) - 1, cmg + 1 + len(data[0][0]) - 1)
        _add_outline_border(ws, rowno, cmg + 1, rowno + len(data[0]) - 1, cmg + 1 + len(data[0][0]) - 1)
        _add_outline_border(ws, rowno, cmg + 1, rowno, cmg + 1 + len(data[0][0]) - 1)
        _add_outline_border(ws, rowno, cmg + 1, rowno + len(data[0]) - 1, cmg + 2)

        # Cabecera (nombres de días) y columnas de hora, resaltadas.
        for col in range(cmg + 1, cmg + 1 + len(data[0][0])):
            ws.cell(row=rowno, column=col).fill = header_fill
        for r_idx in range(rowno + 1, rowno + len(data[0])):
            ws.cell(row=r_idx, column=cmg + 1).fill = header_fill
            ws.cell(row=r_idx, column=cmg + 2).fill = header_fill

    _add_all_borders(ws, fmg + 1, cmg + len(data[0][0]) + 2, fmg + len(summary) + 1, cmg + len(data[0][0]) + 3)
    ws.cell(row=fmg + 1, column=cmg + len(data[0][0]) + 2).fill = header_fill
    ws.cell(row=fmg + 1, column=cmg + len(data[0][0]) + 3).fill = header_fill

    # Hoja "Eventos": lista plana de clases (sin festivos), una fila por hora.
    events = subject_calendar.copy()
    events["length"] = events["end"] - events["start"]
    events = events.reindex(events.index.repeat(events.length)).drop(columns=["length"]).reset_index(drop=True)
    events = events[events.group != holiday]
    events = events.sort_values(["date", "start"])
    events = events[["week", "date", "room", "group"]]
    events.columns = ["Semana", "Fecha", "Aula", "Evento"]

    ws_events = wb.create_sheet(title="Eventos")
    header_font = Font(bold=True)
    events_header_fill = PatternFill(start_color="C0C0C0", end_color="C0C0C0", fill_type="solid")
    for c_idx, header in enumerate(events.columns, 1):
        cell = ws_events.cell(row=2, column=c_idx + 1, value=header)
        cell.font = header_font
        cell.fill = events_header_fill

    for r_idx, row in enumerate(events.values, 2):
        for c_idx, value in enumerate(row, 1):
            ws_events.cell(row=r_idx + 1, column=c_idx + 1, value=str(value) if value is not None else None)

    thin_border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
    for row in ws_events.iter_rows(min_row=2, max_row=len(events) + 2, min_col=2, max_col=len(events.columns) + 1):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")

    ws_events.column_dimensions["C"].width = 11

    result = BytesIO()
    wb.save(result)
    return result.getvalue()
